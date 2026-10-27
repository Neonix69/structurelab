import os, sqlite3, time, re
from pathlib import Path
from flask import Flask, render_template, request, jsonify, session
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import requests

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 15 * 1024 * 1024
app.secret_key = os.environ.get('APP_SECRET') or os.urandom(32)
# Demo accounts are seeded into SQLite once; new accounts persist across restarts.
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / 'data'; DATA_DIR.mkdir(exist_ok=True)
UPLOAD_DIR = DATA_DIR / 'faculty_materials'; UPLOAD_DIR.mkdir(exist_ok=True)
DB_PATH = DATA_DIR / 'learning_world.sqlite3'
ALLOWED_EXTENSIONS = {'.txt', '.md', '.pdf', '.docx', '.pptx'}
def db_conn():
    con = sqlite3.connect(DB_PATH); con.row_factory = sqlite3.Row; return con
def init_db():
    with db_conn() as con:
        con.execute('CREATE TABLE IF NOT EXISTS attempts (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT NOT NULL, topic TEXT NOT NULL, correct INTEGER NOT NULL, difficulty TEXT NOT NULL, created_at TEXT DEFAULT CURRENT_TIMESTAMP)')
        con.execute('CREATE TABLE IF NOT EXISTS materials (id INTEGER PRIMARY KEY AUTOINCREMENT, filename TEXT NOT NULL, stored_name TEXT NOT NULL, extracted_text TEXT NOT NULL, uploaded_by TEXT NOT NULL, approved INTEGER NOT NULL DEFAULT 0, created_at TEXT DEFAULT CURRENT_TIMESTAMP)')
        con.execute("CREATE TABLE IF NOT EXISTS users (username TEXT PRIMARY KEY, display_name TEXT NOT NULL, email TEXT NOT NULL UNIQUE COLLATE NOCASE, password_hash TEXT NOT NULL, role TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'active', created_at TEXT DEFAULT CURRENT_TIMESTAMP)")
        cols = [r['name'] for r in con.execute('PRAGMA table_info(materials)').fetchall()]
        if 'approved' not in cols: con.execute('ALTER TABLE materials ADD COLUMN approved INTEGER NOT NULL DEFAULT 0')
        # Preserve familiar demo logins for evaluation. Passwords are stored as hashes.
        seeds = [
            ('student', 'Student', 'student-demo@local.test', os.environ.get('STUDENT_PASSWORD', 'student123'), 'student', 'active'),
            ('teacher', 'Teacher', 'teacher-demo@local.test', os.environ.get('TEACHER_PASSWORD', 'teacher123'), 'teacher', 'active'),
            ('admin', 'Administrator', 'admin-demo@local.test', os.environ.get('ADMIN_PASSWORD', 'admin123'), 'admin', 'active'),
        ]
        for username, display_name, email, password, role, status in seeds:
            exists = con.execute('SELECT username FROM users WHERE username=?', (username,)).fetchone()
            if not exists:
                con.execute('INSERT INTO users(username,display_name,email,password_hash,role,status) VALUES(?,?,?,?,?,?)', (username, display_name, email, generate_password_hash(password), role, status))
init_db()
# The key lives on the server, never in the student-facing page. For this local demo,
# a teacher sets it once per running server process.
RUNTIME_GROQ_KEY = os.environ.get('GROQ_API_KEY', '').strip()
RUNTIME_MODEL = os.environ.get('GROQ_MODEL', 'openai/gpt-oss-120b')
GROQ_URL = 'https://api.groq.com/openai/v1/chat/completions'
DEFAULT_MODEL = 'openai/gpt-oss-120b'
MODELS = ['openai/gpt-oss-120b', 'openai/gpt-oss-20b']
SYSTEM = '''You are Structure AI, a polished conversational assistant and learning tutor. Respond naturally and directly, with the clarity, warmth, reasoning, and formatting of a high-quality chat assistant. Answer the actual question first. Match the user's tone without fake enthusiasm or repetitive openings. Keep simple answers concise; for complex questions, be thorough and organize the answer into useful sections. Use valid Markdown deliberately: headings for genuinely long answers, paragraphs for explanations, numbered lists for ordered steps, bullets for separate items, Markdown tables for side-by-side comparisons or structured data, and fenced code blocks for code. When using a Markdown table, include a header row and a separator row, e.g. | Item | Meaning | followed by |---|---|. Do not create tables out of repeated dash characters, use long rows of dashes as decoration, or leave raw Markdown syntax unexplained. Avoid excessive bullets when normal prose is clearer. Use bold to highlight important terms sparingly. For code requests, provide complete runnable examples when appropriate, then explain key lines and expected output. For algorithms, explain the idea, trace a concrete example, and include time/space complexity when relevant. For visual requests, provide a clear step-by-step trace and point to the interactive Visual Playground. For hints, reveal help progressively rather than immediately giving away the solution. For unrelated questions, still help without forcing a data-structures connection. Never invent facts; acknowledge uncertainty when needed. Never return an empty response. Be complete, not padded, and do not intentionally truncate useful content.'''

@app.get('/')
def home():
    return render_template('index.html', models=MODELS, default_model=DEFAULT_MODEL)

def logged_in():
    return bool(session.get('role') in ('student', 'teacher', 'admin'))

def can_manage_ai():
    return session.get('role') in ('teacher', 'admin')

@app.get('/api/session')
def session_info():
    return jsonify(ok=True, authenticated=logged_in(), role=session.get('role'), display=session.get('display'), ai_ready=bool(RUNTIME_GROQ_KEY))

@app.post('/api/register')
def register():
    data = request.get_json(silent=True) or {}
    display = ' '.join(str(data.get('full_name') or '').strip().split())[:80]
    username = str(data.get('username') or '').strip().lower()
    email = str(data.get('email') or '').strip().lower()
    password = str(data.get('password') or '')
    requested_role = str(data.get('role') or 'student').strip().lower()
    if not display or not username or not email or not password:
        return jsonify(ok=False, error='Please fill in every field.'), 400
    if not re.fullmatch(r'[a-z0-9_]{3,24}', username):
        return jsonify(ok=False, error='Username must be 3–24 characters using letters, numbers, or underscores.'), 400
    if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email) or len(email) > 254:
        return jsonify(ok=False, error='Enter a valid email address.'), 400
    if len(password) < 8:
        return jsonify(ok=False, error='Use a password with at least 8 characters.'), 400
    if requested_role not in ('student', 'teacher'):
        return jsonify(ok=False, error='Choose Student or Teacher.'), 400
    status = 'active' if requested_role == 'student' else 'pending'
    try:
        with db_conn() as con:
            con.execute('INSERT INTO users(username,display_name,email,password_hash,role,status) VALUES(?,?,?,?,?,?)', (username, display, email, generate_password_hash(password), requested_role, status))
    except sqlite3.IntegrityError:
        return jsonify(ok=False, error='That username or email is already registered. Try signing in or use another one.'), 409
    if requested_role == 'teacher':
        return jsonify(ok=True, pending=True, message='Teacher request submitted. An administrator must approve it before teacher access is enabled.')
    session.clear()
    session['username'] = username
    session['role'] = 'student'
    session['display'] = display
    return jsonify(ok=True, pending=False, role='student', display=display, ai_ready=bool(RUNTIME_GROQ_KEY), message='Account created. Welcome to your learning workspace!')

@app.post('/api/login')
def login():
    data = request.get_json(silent=True) or {}
    username = (data.get('username') or '').strip().lower()
    password = data.get('password') or ''
    with db_conn() as con:
        row = con.execute('SELECT username,display_name,password_hash,role,status FROM users WHERE username=?', (username,)).fetchone()
    if not row or not check_password_hash(row['password_hash'], password):
        return jsonify(ok=False, error='Username or password is incorrect.'), 401
    if row['status'] == 'pending':
        return jsonify(ok=False, error='Your teacher account is awaiting administrator approval.'), 403
    if row['status'] != 'active':
        return jsonify(ok=False, error='This account is not active. Contact the administrator.'), 403
    session.clear()
    session['username'] = row['username']
    session['role'] = row['role']
    session['display'] = row['display_name']
    return jsonify(ok=True, role=row['role'], display=row['display_name'], ai_ready=bool(RUNTIME_GROQ_KEY))

@app.get('/api/admin/teacher-requests')
def admin_teacher_requests():
    if session.get('role') != 'admin': return jsonify(ok=False, error='Administrator access required.'), 403
    with db_conn() as con:
        rows = con.execute("SELECT username,display_name,email,created_at FROM users WHERE role='teacher' AND status='pending' ORDER BY created_at").fetchall()
    return jsonify(ok=True, requests=[dict(username=r['username'], full_name=r['display_name'], email=r['email'], created_at=r['created_at']) for r in rows])

@app.post('/api/admin/teacher-requests/<username>/approve')
def admin_approve_teacher(username):
    if session.get('role') != 'admin': return jsonify(ok=False, error='Administrator access required.'), 403
    with db_conn() as con:
        cur = con.execute("UPDATE users SET status='active' WHERE username=? AND role='teacher' AND status='pending'", (username.lower(),))
        if cur.rowcount == 0: return jsonify(ok=False, error='Pending teacher request not found.'), 404
    return jsonify(ok=True, message=f'Teacher account {username} approved.')

@app.post('/api/logout')
def logout():
    session.clear()
    return jsonify(ok=True)

@app.post('/api/test')
def test_connection():
    global RUNTIME_GROQ_KEY, RUNTIME_MODEL
    if not can_manage_ai():
        return jsonify(ok=False, error='Only teachers or administrators can configure the AI connection.')
    data = request.get_json(silent=True) or {}
    key = (data.get('api_key') or '').strip()
    model = data.get('model') or DEFAULT_MODEL
    if not key:
        return jsonify(ok=False, error='Paste your Groq API key first.')
    result = call_groq(key, model, [{'role':'user','content':'Reply exactly: Connection successful.'}], max_tokens=40)
    if result['ok']:
        RUNTIME_GROQ_KEY = key
        RUNTIME_MODEL = model
        return jsonify(ok=True, message='Groq connected successfully for this running application.', model=model)
    return jsonify(ok=False, error=result['error'])

@app.post('/api/chat')
def chat():
    if not logged_in():
        return jsonify(ok=False, error='Please sign in to use the AI Tutor.')
    data = request.get_json(silent=True) or {}
    key = RUNTIME_GROQ_KEY
    model = RUNTIME_MODEL or DEFAULT_MODEL
    prompt = (data.get('message') or '').strip()
    if not prompt:
        return jsonify(ok=False, error='Type a question before sending.')
    if not key:
        return jsonify(ok=False, error='The AI service is not configured yet. Please ask your teacher to connect Groq.')
    level = data.get('level', 'College / Advanced')
    concept = data.get('concept', 'General')
    mode = data.get('mode', 'Explain')
    history = data.get('history') or []
    materials = retrieve_materials(prompt, limit=3)
    context_text = ('\n\nFACULTY-APPROVED MATERIAL EXCERPTS (use when relevant; distinguish from general explanation):\n' + '\n---\n'.join(materials)) if materials else ''
    messages = [{'role':'system','content':SYSTEM + f'\nCurrent topic: {concept}. Learner level: {level}. Learning mode: {mode}.' + context_text}]
    for item in history[-12:]:
        role = item.get('role')
        content = item.get('content')
        if role in ('user','assistant') and isinstance(content,str) and content.strip():
            messages.append({'role':role,'content':content[:7000]})
    messages.append({'role':'user','content':prompt})
    result = call_groq(key, model, messages, max_tokens=5000)
    if result['ok']:
        return jsonify(ok=True, answer=result['answer'], model=model)
    return jsonify(ok=False, error=result['error'])



def retrieve_materials(prompt, limit=3):
    words = set(re.findall(r"[a-zA-Z0-9_]{3,}", prompt.lower()))
    if not words: return []
    with db_conn() as con:
        rows = con.execute('SELECT filename, extracted_text FROM materials WHERE approved=1 ORDER BY id DESC').fetchall()
    scored = []
    for row in rows:
        text = row['extracted_text'] or ''
        lower = text.lower()
        score = sum(lower.count(w) for w in words)
        if score: scored.append((score, row['filename'], text))
    scored.sort(reverse=True, key=lambda x: x[0])
    return [f"Source: {name}\n{text[:3500]}" for _, name, text in scored[:limit]]

def extract_text(path, ext):
    try:
        if ext in ('.txt', '.md'):
            return path.read_text(encoding='utf-8', errors='replace')[:100000]
        if ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(str(path)); return '\n'.join((page.extract_text() or '') for page in reader.pages[:80])[:100000]
        if ext == '.docx':
            from docx import Document
            doc = Document(str(path)); return '\n'.join(p.text for p in doc.paragraphs)[:100000]
        if ext == '.pptx':
            from pptx import Presentation
            prs = Presentation(str(path)); return '\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if getattr(shape, 'has_text_frame', False))[:100000]
    except Exception as exc:
        raise ValueError(f'Could not extract text from this file: {exc}')
    raise ValueError('Unsupported file type. Upload TXT, MD, PDF, DOCX, or PPTX.')

@app.get('/api/progress')
def get_progress():
    if not logged_in(): return jsonify(ok=False, error='Please sign in.'), 401
    username = session['username']
    with db_conn() as con:
        rows = con.execute('SELECT topic, correct, difficulty, created_at FROM attempts WHERE username=? ORDER BY id DESC LIMIT 300', (username,)).fetchall()
    attempts = [dict(topic=r['topic'], correct=bool(r['correct']), difficulty=r['difficulty'], created_at=r['created_at']) for r in rows]
    return jsonify(ok=True, attempts=attempts)

@app.post('/api/progress')
def save_progress():
    if not logged_in(): return jsonify(ok=False, error='Please sign in.'), 401
    data = request.get_json(silent=True) or {}
    topic = str(data.get('topic') or 'General')[:80]
    difficulty = str(data.get('difficulty') or 'Beginner')[:30]
    correct = 1 if data.get('correct') else 0
    with db_conn() as con:
        con.execute('INSERT INTO attempts(username, topic, correct, difficulty) VALUES(?,?,?,?)', (session['username'], topic, correct, difficulty))
    return jsonify(ok=True)

@app.get('/api/teacher/analytics')
def teacher_analytics():
    if not can_manage_ai(): return jsonify(ok=False, error='Teacher access required.'), 403
    with db_conn() as con:
        total = con.execute('SELECT COUNT(*) FROM attempts').fetchone()[0]
        accuracy = con.execute('SELECT AVG(correct)*100 FROM attempts').fetchone()[0] or 0
        topics = con.execute('SELECT topic, COUNT(*) n, AVG(correct)*100 accuracy FROM attempts GROUP BY topic ORDER BY n DESC').fetchall()
        materials = con.execute('SELECT id, filename, uploaded_by, approved, created_at FROM materials ORDER BY id DESC LIMIT 20').fetchall()
    return jsonify(ok=True, total_attempts=total, accuracy=round(accuracy), topics=[dict(topic=r['topic'], attempts=r['n'], accuracy=round(r['accuracy'] or 0)) for r in topics], materials=[dict(id=r['id'], filename=r['filename'], uploaded_by=r['uploaded_by'], approved=bool(r['approved']), created_at=r['created_at']) for r in materials])

@app.post('/api/teacher/material/<int:material_id>/approve')
def approve_material(material_id):
    if not can_manage_ai(): return jsonify(ok=False, error='Teacher access required.'), 403
    data = request.get_json(silent=True) or {}
    approved = 1 if data.get('approved', True) else 0
    with db_conn() as con:
        cur = con.execute('UPDATE materials SET approved=? WHERE id=?', (approved, material_id))
        if cur.rowcount == 0: return jsonify(ok=False, error='Material not found.'), 404
    return jsonify(ok=True, message='Material approved and available to the tutor.' if approved else 'Material returned to draft.')

@app.post('/api/teacher/upload')
def teacher_upload():
    if not can_manage_ai(): return jsonify(ok=False, error='Teacher access required.'), 403
    if 'file' not in request.files: return jsonify(ok=False, error='Choose a file to upload.'), 400
    file = request.files['file']
    filename = secure_filename(file.filename or '')
    ext = Path(filename).suffix.lower()
    if not filename or ext not in ALLOWED_EXTENSIONS: return jsonify(ok=False, error='Supported formats: TXT, MD, PDF, DOCX, PPTX.'), 400
    stored = f"{int(time.time())}_{filename}"; path = UPLOAD_DIR / stored
    file.save(path)
    try: text = extract_text(path, ext)
    except ValueError as exc:
        path.unlink(missing_ok=True); return jsonify(ok=False, error=str(exc)), 400
    if not text.strip():
        path.unlink(missing_ok=True); return jsonify(ok=False, error='No readable text was extracted. Scanned PDFs may need OCR before upload.'), 400
    with db_conn() as con:
        con.execute('INSERT INTO materials(filename, stored_name, extracted_text, uploaded_by, approved) VALUES(?,?,?,?,0)', (filename, stored, text, session.get('username','teacher')))
    return jsonify(ok=True, message=f'Uploaded {filename} as a draft. Approve it before the tutor can use its content.', filename=filename)

def call_groq(key, model, messages, max_tokens=5000):
    try:
        r = requests.post(GROQ_URL, headers={'Authorization':f'Bearer {key}','Content-Type':'application/json'}, json={'model':model,'messages':messages,'temperature':0.35,'max_tokens':max_tokens}, timeout=50)
        if not r.ok:
            try:
                body = r.json(); detail = body.get('error',{}).get('message') or r.text[:350]
            except ValueError: detail = r.text[:350]
            if r.status_code == 401: detail = 'Groq rejected this API key (401). Check that it is copied correctly and is still active.'
            elif r.status_code == 403: detail = 'Groq denied this request (403). Check your account access and model permissions.'
            elif r.status_code == 404: detail = f'Model not found (404). Choose another model in AI Settings. Details: {detail}'
            elif r.status_code == 429: detail = 'Groq rate limit or quota reached (429). Wait and retry, or check your Groq Console limits.'
            return {'ok':False,'error':f'Groq request failed ({r.status_code}): {detail}'}
        data = r.json()
        answer = (((data.get('choices') or [{}])[0].get('message') or {}).get('content') or '').strip()
        if not answer: return {'ok':False,'error':'Groq returned an empty message. Try again or choose a different model.'}
        return {'ok':True,'answer':answer}
    except requests.Timeout:
        return {'ok':False,'error':'Groq took too long to respond. Please retry.'}
    except requests.RequestException as exc:
        return {'ok':False,'error':f'Could not reach Groq. Check your internet connection. Details: {exc}'}

if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5000, debug=False)
